"""Tests for mathtrans.export (bilingual PDF, DOCX via LibreOffice and python-docx)."""
from __future__ import annotations

import os
import stat
import time
import unicodedata
from pathlib import Path

import docx
import pymupdf
import pytest
from docx.oxml.ns import qn

from mathtrans import export as export_module
from mathtrans.export import (BILINGUAL_GAP, ExportError, build_docx_from_segments,
                              convert_pdf_to_docx_libreoffice, export_docx, libreoffice_executable,
                              make_bilingual_pdf)
from mathtrans.layout import render_document
from mathtrans.models import Lang
from mathtrans.samples import sample_texts
from test_layout import LONG_EX3, build_document, page_text, sample_translations


def _norm(text: str) -> str:
    return "".join(unicodedata.normalize("NFKC", text).split())


def docx_text(path: Path) -> str:
    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return _norm("\n".join(parts))


@pytest.fixture(scope="module")
def zh_en(sample_pdf_zh, tmp_path_factory):
    translations = sample_translations(Lang.ZH, Lang.EN, {"ex3": LONG_EX3})
    doc = build_document(sample_pdf_zh, Lang.ZH, Lang.EN, translations)
    doc.title = "Chapter 1 The Pythagorean Theorem"
    out = tmp_path_factory.mktemp("export") / "zh_en.pdf"
    render_document(sample_pdf_zh, doc, out)
    return doc, out


@pytest.fixture(scope="module")
def en_zh(sample_pdf_en, tmp_path_factory):
    translations = sample_translations(Lang.EN, Lang.ZH)
    doc = build_document(sample_pdf_en, Lang.EN, Lang.ZH, translations)
    out = tmp_path_factory.mktemp("export_zh") / "en_zh.pdf"
    render_document(sample_pdf_en, doc, out)
    return doc, out


# --------------------------------------------------------------------------- #
# bilingual PDF
# --------------------------------------------------------------------------- #


def test_bilingual_pdf_side_by_side(sample_pdf_zh, zh_en, tmp_path):
    _doc, translated = zh_en
    out = make_bilingual_pdf(sample_pdf_zh, translated, tmp_path / "nested" / "bilingual.pdf")
    assert Path(out).is_file()
    src = pymupdf.open(str(sample_pdf_zh))
    bil = pymupdf.open(out)
    assert bil.page_count == src.page_count == 2
    for pno in range(src.page_count):
        w, h = src[pno].rect.width, src[pno].rect.height
        page = bil[pno]
        assert page.rect.width == pytest.approx(2 * w + BILINGUAL_GAP) and BILINGUAL_GAP == 10
        assert page.rect.height == pytest.approx(h)
        left = page_text(page, clip=pymupdf.Rect(0, 0, w, h))
        right = page_text(page, clip=pymupdf.Rect(w + BILINGUAL_GAP, 0, 2 * w + BILINGUAL_GAP, h))
        assert "勾股定理" in left and "Pythagorean" not in left
        assert _norm("Pythagorean") in right and "勾股定理" not in right
        assert len(page.get_image_info()) == 2 * len(src[pno].get_image_info())


def test_bilingual_pdf_tolerates_page_count_mismatch(sample_pdf_zh, zh_en, tmp_path):
    _doc, translated = zh_en
    one_page = pymupdf.open(str(translated))
    one_page.delete_page(1)
    short = tmp_path / "short.pdf"
    one_page.save(str(short))
    out = make_bilingual_pdf(sample_pdf_zh, short, tmp_path / "bilingual.pdf")
    bil = pymupdf.open(out)
    assert bil.page_count == 2
    w = pymupdf.open(str(sample_pdf_zh))[1].rect.width
    assert "练习" in page_text(bil[1], clip=pymupdf.Rect(0, 0, w, 842))
    assert page_text(bil[1], clip=pymupdf.Rect(w + BILINGUAL_GAP, 0, 2 * w + BILINGUAL_GAP, 842)) == ""


# --------------------------------------------------------------------------- #
# DOCX: python-docx builder
# --------------------------------------------------------------------------- #


def test_export_docx_python_docx_from_segments(zh_en, tmp_path):
    doc, translated = zh_en
    out = export_docx(translated, tmp_path / "out" / "translated.docx", doc, use_libreoffice=False)
    assert Path(out).suffix == ".docx" and Path(out).is_file()
    document = docx.Document(out)
    text = docx_text(Path(out))
    en = sample_texts("en")
    for key in ("title", "para1", "solution", "ex1", "think"):
        assert _norm(en[key]) in text
    assert _norm(LONG_EX3) in text and "勾股定理" not in text
    assert _norm("a² + b² = c²") in text  # untranslated formula is kept
    headings = [p for p in document.paragraphs if p.style.name.startswith("Heading")]
    assert [p.text for p in headings][:2] == ["Chapter 1 The Pythagorean Theorem", "1.1 Exploring the Pythagorean Theorem"]
    assert any(p.text == "Exercises 1.1" for p in headings)
    assert len(document.inline_shapes) == 2  # the triangle and the square figure
    assert document.core_properties.title == "Chapter 1 The Pythagorean Theorem"
    # reading order: the figure follows its paragraph and precedes the caption on page 1
    kinds = []
    for p in document.paragraphs:
        if p._p.xpath(".//pic:pic"):
            kinds.append("image")
        elif p.text.strip():
            kinds.append(p.text[:12])
    assert kinds.index("image") == kinds.index("In a right t") + 1
    assert kinds.index("Figure 1-1 S") == kinds.index("image") + 1


def test_export_docx_python_docx_without_segments(zh_en, tmp_path):
    _doc, translated = zh_en
    out = build_docx_from_segments(translated, tmp_path / "blocks.docx", doc=None)
    text = docx_text(Path(out))
    assert _norm(sample_texts("en")["theorem"]) in text and _norm("Exercises 1.1") in text
    assert len(docx.Document(out).inline_shapes) == 2


def test_export_docx_cjk_target_sets_east_asian_font(en_zh, tmp_path):
    doc, translated = en_zh
    out = export_docx(translated, tmp_path / "zh.docx", doc, use_libreoffice=False)
    text = docx_text(Path(out))
    zh = sample_texts("zh")
    assert _norm(zh["title"]) in text and _norm(zh["ex3"]) in text and "Pythagorean" not in text
    document = docx.Document(out)
    rfonts = document.styles["Normal"].element.rPr.find(qn("w:rFonts"))
    assert rfonts.get(qn("w:eastAsia")) == "Microsoft YaHei"


def test_export_docx_auto_falls_back_without_libreoffice(zh_en, tmp_path, monkeypatch):
    doc, translated = zh_en
    monkeypatch.setattr(export_module, "libreoffice_executable", lambda: None)
    out = export_docx(translated, tmp_path / "auto.docx", doc)
    assert Path(out).is_file() and _norm("Exercises 1.1") in docx_text(Path(out))
    with pytest.raises(ExportError):
        export_docx(translated, tmp_path / "forced.docx", doc, use_libreoffice=True)


def test_export_docx_auto_falls_back_when_conversion_fails(zh_en, tmp_path, monkeypatch):
    doc, translated = zh_en
    fake = tmp_path / "soffice"
    fake.write_text("#!/bin/sh\necho 'Error: source file could not be loaded' >&2\nexit 0\n")
    fake.chmod(fake.stat().st_mode | stat.S_IXUSR)
    monkeypatch.setattr(export_module, "libreoffice_executable", lambda: str(fake))
    out = export_docx(translated, tmp_path / "auto.docx", doc)
    assert _norm("Exercises 1.1") in docx_text(Path(out))
    with pytest.raises(ExportError, match="conversion failed"):
        export_docx(translated, tmp_path / "forced.docx", doc, use_libreoffice=True)


def test_libreoffice_timeout_kills_the_process(zh_en, tmp_path, monkeypatch):
    _doc, translated = zh_en
    fake = tmp_path / "soffice"
    fake.write_text("#!/bin/sh\nsleep 60\n")
    fake.chmod(fake.stat().st_mode | stat.S_IXUSR)
    monkeypatch.setattr(export_module, "libreoffice_executable", lambda: str(fake))
    started = time.monotonic()
    with pytest.raises(ExportError, match="did not finish"):
        convert_pdf_to_docx_libreoffice(translated, tmp_path / "timeout.docx", timeout=1)
    assert time.monotonic() - started < 15
    assert not (tmp_path / "timeout.docx").exists()


# --------------------------------------------------------------------------- #
# DOCX: LibreOffice (skipped when it is unavailable or cannot import PDF here)
# --------------------------------------------------------------------------- #


def test_export_docx_with_libreoffice(zh_en, tmp_path):
    _doc, translated = zh_en
    if libreoffice_executable() is None:
        pytest.skip("LibreOffice (soffice) is not installed")
    try:
        out = convert_pdf_to_docx_libreoffice(translated, tmp_path / "lo.docx", timeout=180)
    except ExportError as exc:
        pytest.skip(f"LibreOffice cannot convert PDF to DOCX here: {exc}")
    assert Path(out).is_file()
    text = docx_text(Path(out))
    assert _norm("Pythagorean") in text
    assert os.path.getsize(out) > 1000


# --------------------------------------------------------------------------- #
# adversarial: rotated / odd pages, signature compatibility, images, ligatures
# --------------------------------------------------------------------------- #

CSS_PLAIN = "* {font-family: sans-serif;} p {margin:0}"


def _pixel_mismatch(a: pymupdf.Pixmap, b: pymupdf.Pixmap) -> float:
    """Share of differing bytes between two same-sized pixmaps (0.0 = identical)."""
    assert (a.width, a.height, a.n) == (b.width, b.height, b.n)
    sa, sb = a.samples, b.samples
    return sum(1 for x, y in zip(sa, sb) if x != y) / max(len(sa), 1)


def test_bilingual_pdf_honours_page_rotation(tmp_path):
    """show_pdf_page ignores /Rotate; the bilingual page must still show each page
    the way a viewer displays it (checked pixel-wise against the page render)."""
    src = pymupdf.open()
    for rotation, label in ((90, "ninety"), (270, "two-seventy"), (180, "upside")):
        page = src.new_page(width=300, height=150)
        page.insert_htmlbox(pymupdf.Rect(10, 10, 200, 40), f'<p style="font-size:16px">{label}</p>', css=CSS_PLAIN)
        page.draw_rect(pymupdf.Rect(220, 100, 290, 140), color=(0, 0, 1), fill=(0.8, 0.9, 1), width=2)
        page.set_rotation(rotation)
    src_path = tmp_path / "rotated_src.pdf"
    src.save(str(src_path))
    out = make_bilingual_pdf(src_path, src_path, tmp_path / "bilingual.pdf")
    bil = pymupdf.open(out)
    src = pymupdf.open(str(src_path))
    assert bil.page_count == 3
    for pno in range(3):
        w, h = src[pno].rect.width, src[pno].rect.height  # displayed size (150 x 300 for 90 / 270)
        assert (w, h) == ((150, 300) if pno < 2 else (300, 150))
        page = bil[pno]
        assert page.rect.width == pytest.approx(2 * w + BILINGUAL_GAP) and page.rect.height == pytest.approx(h)
        expected = src[pno].get_pixmap(dpi=72)
        left = page.get_pixmap(dpi=72, clip=pymupdf.Rect(0, 0, w, h))
        right = page.get_pixmap(dpi=72, clip=pymupdf.Rect(w + BILINGUAL_GAP, 0, 2 * w + BILINGUAL_GAP, h))
        assert _pixel_mismatch(expected, left) < 0.01, f"page {pno}: left half is not the displayed source page"
        assert _pixel_mismatch(expected, right) < 0.01, f"page {pno}: right half is not the displayed page"
        assert _pixel_mismatch(expected, pymupdf.Pixmap(expected.colorspace, expected.irect, 0)) > 0.02  # not blank


def test_bilingual_pdf_with_different_page_sizes(sample_pdf_zh, tmp_path):
    big = pymupdf.open()
    page = big.new_page(width=800, height=1000)
    page.insert_htmlbox(pymupdf.Rect(50, 50, 700, 100), '<p style="font-size:20px">big translated page</p>', css=CSS_PLAIN)
    big.new_page(width=200, height=300)
    big_path = tmp_path / "big.pdf"
    big.save(str(big_path))
    out = make_bilingual_pdf(sample_pdf_zh, big_path, tmp_path / "mixed.pdf")
    bil = pymupdf.open(out)
    assert bil[0].rect.width == pytest.approx(595 + BILINGUAL_GAP + 800) and bil[0].rect.height == pytest.approx(1000)
    assert bil[1].rect.width == pytest.approx(595 + BILINGUAL_GAP + 200) and bil[1].rect.height == pytest.approx(842)
    assert "勾股定理" in page_text(bil[0], clip=pymupdf.Rect(0, 0, 595, 842))
    assert _norm("big translated page") in page_text(bil[0], clip=pymupdf.Rect(605, 0, 1405, 1000))
    assert page_text(bil[0], clip=pymupdf.Rect(0, 842, 595, 1000)) == ""  # below the shorter page: blank


def test_bilingual_pdf_rejects_documents_without_pages(tmp_path):
    empty = tmp_path / "empty.pdf"
    empty.write_bytes(b"%PDF-1.4\n1 0 obj\n<< /Type /Catalog /Pages 2 0 R >>\nendobj\n"
                      b"2 0 obj\n<< /Type /Pages /Kids [] /Count 0 >>\nendobj\ntrailer\n<< /Root 1 0 R >>\n%%EOF\n")
    assert pymupdf.open(str(empty)).page_count == 0
    with pytest.raises(ExportError, match="has pages"):
        make_bilingual_pdf(empty, empty, tmp_path / "out.pdf")
    assert not (tmp_path / "out.pdf").exists()


def test_export_docx_contract_signature(zh_en, tmp_path, monkeypatch):
    """ARCHITECTURE: export_docx(translated_pdf, out_path, doc=None, timeout=180) - the
    timeout must be the 4th positional parameter; the LibreOffice switch is keyword-only."""
    doc, translated = zh_en
    monkeypatch.setattr(export_module, "libreoffice_executable", lambda: None)
    out = export_docx(translated, tmp_path / "positional.docx", doc, 5)
    assert _norm("Exercises 1.1") in docx_text(Path(out))
    with pytest.raises(TypeError):
        export_docx(translated, tmp_path / "x.docx", doc, 5, True)  # type: ignore[misc]
    out = export_docx(translated, tmp_path / "nodoc.docx", None, 5)  # doc=None: blocks path
    assert _norm("Exercises 1.1") in docx_text(Path(out))


def _docx_image_blobs(path: Path) -> list[bytes]:
    document = docx.Document(str(path))
    blobs = []
    for shape in document.inline_shapes:
        rid = shape._inline.graphic.graphicData.pic.blipFill.blip.embed
        blobs.append(document.part.related_parts[rid].blob)
    return blobs


def test_docx_keeps_image_transparency_and_handles_cmyk(tmp_path):
    import io

    from PIL import Image

    rgba = Image.new("RGBA", (40, 40), (255, 0, 0, 0))
    for x in range(10, 30):
        for y in range(10, 30):
            rgba.putpixel((x, y), (0, 0, 255, 255))
    buf_rgba = io.BytesIO()
    rgba.save(buf_rgba, format="PNG")
    cmyk = Image.new("CMYK", (30, 20), (0, 255, 255, 0))
    buf_cmyk = io.BytesIO()
    cmyk.save(buf_cmyk, format="JPEG")
    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_htmlbox(pymupdf.Rect(50, 40, 500, 70), '<p style="font-size:12px">Two images follow</p>', css=CSS_PLAIN)
    page.insert_image(pymupdf.Rect(50, 100, 150, 200), stream=buf_rgba.getvalue())
    page.insert_image(pymupdf.Rect(50, 250, 200, 350), stream=buf_cmyk.getvalue())
    src = tmp_path / "images.pdf"
    pdf.save(str(src))
    infos = pymupdf.open(str(src)).extract_image(pymupdf.open(str(src))[0].get_image_info(xrefs=True)[0]["xref"])
    assert infos["smask"] > 0  # the PNG alpha became a soft mask
    out = build_docx_from_segments(src, tmp_path / "images.docx", doc=None)
    blobs = _docx_image_blobs(Path(out))
    assert len(blobs) == 2
    with Image.open(io.BytesIO(blobs[0])) as im:
        assert im.mode == "RGBA" and im.getpixel((0, 0))[3] == 0 and im.getpixel((20, 20)) == (0, 0, 255, 255)
    with Image.open(io.BytesIO(blobs[1])) as im:
        assert im.mode in ("RGB", "L") and im.size == (30, 20)  # CMYK converted for Word
    assert _norm("Two images follow") in docx_text(Path(out))


def test_docx_blocks_path_expands_ligatures(tmp_path):
    pdf = pymupdf.open()
    pdf.new_page().insert_htmlbox(pymupdf.Rect(50, 40, 500, 70),
                                  '<p style="font-size:12px">office fluffy waffle</p>', css=CSS_PLAIN)
    src = tmp_path / "liga.pdf"
    pdf.save(str(src))
    assert "ﬃ" in pymupdf.open(str(src))[0].get_text()  # the PDF really contains ligature glyphs
    out = build_docx_from_segments(src, tmp_path / "liga.docx", doc=None)
    raw = "\n".join(p.text for p in docx.Document(out).paragraphs)
    assert "office fluffy waffle" in raw and not any(0xFB00 <= ord(c) <= 0xFB06 for c in raw)
