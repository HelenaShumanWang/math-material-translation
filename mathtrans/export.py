"""Exports derived from the translated PDF: bilingual side-by-side PDF and DOCX.

* :func:`make_bilingual_pdf` puts every original page next to its translated
  page (``page.show_pdf_page``) for proofreading.
* :func:`export_docx` converts the translated PDF with LibreOffice
  (``writer_pdf_import``) when it is installed and works, and otherwise builds a
  DOCX with python-docx from the document's segments (headings, paragraphs and
  the page images in reading order), so a file is always produced.
"""
from __future__ import annotations

import io
import logging
import os
import shutil
import signal
import subprocess
import tempfile
from pathlib import Path
from typing import Optional, Union

import docx
import pymupdf
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from .languages import detect_language, is_cjk
from .models import BBox, Lang, SegmentKind, TextSegment, TranslatedDocument

log = logging.getLogger("mathtrans.export")

PathLike = Union[str, "os.PathLike[str]"]

BILINGUAL_GAP = 10.0
"""Points between the original and the translated page in the bilingual PDF."""
MAX_IMAGE_WIDTH_IN = 6.0
"""Widest picture inserted in the fallback DOCX (fits a portrait page with margins)."""
_LIBREOFFICE_NAMES = ("soffice", "libreoffice")
_EAST_ASIAN_FONTS = {Lang.ZH: "Microsoft YaHei", Lang.JA: "Yu Gothic", Lang.KO: "Malgun Gothic"}
_ALIGNMENTS = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}


class ExportError(RuntimeError):
    """An export step could not produce its file."""


# --------------------------------------------------------------------------- #
# Bilingual PDF
# --------------------------------------------------------------------------- #


def _show_page(target: pymupdf.Page, rect: pymupdf.Rect, source: pymupdf.Document, pno: int) -> None:
    """Place source page ``pno`` into ``rect`` as it is *displayed*.

    ``show_pdf_page`` draws the unrotated page content but derives the source
    rectangle from the rotated ``page.rect``, so a page with ``/Rotate`` would
    come out unrotated and scaled down. The rotation is therefore zeroed on the
    in-memory page (``source`` is never saved), applied explicitly (PyMuPDF's
    ``rotate`` turns counter-clockwise, ``/Rotate`` clockwise) and restored.
    """
    page = source[pno]
    rotation = page.rotation
    if not rotation:
        target.show_pdf_page(rect, source, pno)
        return
    page.set_rotation(0)
    try:
        target.show_pdf_page(rect, source, pno, rotate=-rotation)
    finally:
        page.set_rotation(rotation)


def make_bilingual_pdf(src_pdf: PathLike, translated_pdf: PathLike, out_path: PathLike,
                       gap: float = BILINGUAL_GAP) -> str:
    """Write a PDF whose page *i* shows source page *i* on the left and the
    translated page *i* on the right (``width = w_src + gap + w_translated``,
    height = the taller of the two). A missing page on either side leaves that
    half blank; pages are shown as displayed (``/Rotate`` honoured).
    Raises :class:`ExportError` when neither input has a page.
    """
    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    with pymupdf.open(str(src_pdf)) as src, pymupdf.open(str(translated_pdf)) as tr, pymupdf.open() as result:
        if src.page_count != tr.page_count:
            log.warning("bilingual export: page counts differ (source %d, translation %d)",
                        src.page_count, tr.page_count)
        n_pages = max(src.page_count, tr.page_count)
        if n_pages == 0:
            raise ExportError(f"cannot build a bilingual PDF: neither {src_pdf} nor {translated_pdf} has pages")
        for pno in range(n_pages):
            left = src[pno].rect if pno < src.page_count else tr[pno].rect
            right = tr[pno].rect if pno < tr.page_count else src[pno].rect
            width = left.width + gap + right.width
            height = max(left.height, right.height)
            page = result.new_page(width=width, height=height)
            if pno < src.page_count:
                _show_page(page, pymupdf.Rect(0, 0, left.width, left.height), src, pno)
            if pno < tr.page_count:
                _show_page(page, pymupdf.Rect(left.width + gap, 0, width, right.height), tr, pno)
            x = left.width + gap / 2
            page.draw_line(pymupdf.Point(x, 0), pymupdf.Point(x, height), color=(0.75, 0.75, 0.75), width=0.5)
        result.save(str(out), garbage=4, deflate=True)
    log.info("wrote bilingual PDF %s (%d pages)", out, n_pages)
    return str(out)


# --------------------------------------------------------------------------- #
# DOCX via LibreOffice
# --------------------------------------------------------------------------- #


def libreoffice_executable() -> Optional[str]:
    """Path of the LibreOffice launcher, or None when it is not installed."""
    for name in _LIBREOFFICE_NAMES:
        path = shutil.which(name)
        if path:
            return path
    return None


def convert_pdf_to_docx_libreoffice(pdf_path: PathLike, out_path: PathLike, timeout: float = 180) -> str:
    """Convert with ``soffice --headless --infilter=writer_pdf_import --convert-to docx``.

    Runs with a private ``HOME``/user profile (no profile-lock clashes with a
    running LibreOffice), kills the whole process group on timeout and raises
    :class:`ExportError` whenever no valid DOCX comes out.
    """
    exe = libreoffice_executable()
    if exe is None:
        raise ExportError("LibreOffice (soffice) is not installed or not on PATH")
    pdf = Path(pdf_path).resolve()
    if not pdf.is_file():
        raise ExportError(f"PDF not found: {pdf}")
    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="mathtrans-lo-") as tmp:
        home = Path(tmp) / "home"
        profile = Path(tmp) / "profile"
        outdir = Path(tmp) / "out"
        for d in (home, profile, outdir):
            d.mkdir()
        env = dict(os.environ, HOME=str(home))
        cmd = [exe, "--headless", "--norestore", "--nologo", "--nodefault",
               f"-env:UserInstallation={profile.as_uri()}",
               "--infilter=writer_pdf_import", "--convert-to", "docx", "--outdir", str(outdir), str(pdf)]
        log.info("running LibreOffice: %s", " ".join(cmd))
        proc = subprocess.Popen(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, env=env,
                                start_new_session=True)
        try:
            stdout, stderr = proc.communicate(timeout=timeout)
        except subprocess.TimeoutExpired:
            _kill_process_group(proc)
            raise ExportError(f"LibreOffice did not finish within {timeout:.0f}s") from None
        produced = outdir / (pdf.stem + ".docx")
        if proc.returncode != 0 or not produced.is_file():
            detail = (stderr or stdout or "").strip().splitlines()
            raise ExportError("LibreOffice conversion failed (exit code %s): %s"
                              % (proc.returncode, detail[-1] if detail else "no output"))
        try:
            docx.Document(str(produced))
        except Exception as exc:
            raise ExportError(f"LibreOffice produced an unreadable DOCX: {exc}") from exc
        shutil.move(str(produced), str(out))
    log.info("LibreOffice wrote %s", out)
    return str(out)


def _kill_process_group(proc: subprocess.Popen) -> None:
    try:
        os.killpg(os.getpgid(proc.pid), signal.SIGKILL)
    except (ProcessLookupError, PermissionError):
        proc.kill()
    try:
        proc.communicate(timeout=10)
    except subprocess.TimeoutExpired:  # pragma: no cover - the group was killed with SIGKILL
        log.warning("LibreOffice process %d did not terminate after SIGKILL", proc.pid)


# --------------------------------------------------------------------------- #
# DOCX via python-docx (fallback)
# --------------------------------------------------------------------------- #


def _configure_fonts(document: docx.document.Document, lang: Optional[Lang]) -> None:
    style = document.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(11)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), _EAST_ASIAN_FONTS.get(lang, "Microsoft YaHei"))


def _image_bytes(pdf: pymupdf.Document, xref: int) -> Optional[bytes]:
    """PNG/JPEG bytes of image ``xref`` in a form python-docx accepts (None if unusable).

    PNG / JPEG streams without a soft mask are passed through; everything else
    (JPEG 2000, CCITT, CMYK, masked images ...) is rendered to PNG, with the
    soft mask composited into an alpha channel so transparency survives.
    """
    try:
        info = pdf.extract_image(xref)
    except Exception as exc:
        log.warning("cannot extract image xref %d: %s", xref, exc)
        return None
    smask = int(info.get("smask") or 0) if info else 0
    if info and info.get("ext") in ("png", "jpeg", "jpg") and not smask:
        return info["image"]
    try:
        pix = pymupdf.Pixmap(pdf, xref)
        if pix.n - pix.alpha >= 4:  # CMYK and friends
            pix = pymupdf.Pixmap(pymupdf.csRGB, pix)
        if smask and not pix.alpha:
            mask = pymupdf.Pixmap(pdf, smask)
            if (mask.width, mask.height) == (pix.width, pix.height) and mask.n == 1:
                pix = pymupdf.Pixmap(pix, mask)
            else:
                log.warning("image xref %d: soft mask %d has another size; transparency dropped", xref, smask)
        return pix.tobytes("png")
    except Exception as exc:
        log.warning("cannot convert image xref %d to PNG: %s", xref, exc)
        return None


def _add_picture(document: docx.document.Document, data: bytes, bbox: BBox) -> bool:
    width_in = bbox.width / 72.0 if bbox.width > 0 else MAX_IMAGE_WIDTH_IN
    width_in = max(0.5, min(width_in, MAX_IMAGE_WIDTH_IN))
    try:
        document.add_picture(io.BytesIO(data), width=Inches(width_in))
    except Exception as exc:
        log.warning("python-docx rejected an image (%d bytes): %s", len(data), exc)
        return False
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    return True


def _add_segment(document: docx.document.Document, seg: TextSegment) -> None:
    text = seg.effective_text.strip()
    if not text:
        return
    role = seg.style.role
    if role == "heading":
        level = 1 if seg.style.size >= 16 else 2
        document.add_heading(text, level=level)
        return
    paragraph = document.add_paragraph()
    paragraph.alignment = _ALIGNMENTS.get(seg.style.align, WD_ALIGN_PARAGRAPH.LEFT)
    run = paragraph.add_run(text)
    run.bold = bool(seg.style.bold)
    run.italic = bool(seg.style.italic or role == "caption")
    if seg.style.size > 0:
        run.font.size = Pt(max(6.0, min(seg.style.size, 36.0)))


def _page_images(page: pymupdf.Page) -> list[tuple[BBox, int]]:
    """(placement bbox, xref) for every image placed on the page, in page order."""
    page_rect = BBox.from_rect(page.rect)
    seen: set[tuple[int, tuple[float, ...]]] = set()
    out: list[tuple[BBox, int]] = []
    for info in page.get_image_info(xrefs=True):
        xref = int(info.get("xref") or 0)
        if xref <= 0:
            continue
        bbox = BBox.from_rect(info["bbox"])
        if bbox.intersection_area(page_rect) <= 0:
            continue
        key = (xref, tuple(round(v, 1) for v in bbox.as_tuple()))
        if key in seen:
            continue
        seen.add(key)
        out.append((bbox, xref))
    return out


def _image_slots(segments: list[TextSegment], images: list[tuple[BBox, int]]) -> tuple[dict[int, list], list]:
    """Decide after which segment (index) each image is emitted.

    An image follows the nearest segment that ends above it in the same column
    (horizontal overlap); failing that it precedes the first segment below it
    in that column, then the first segment starting below its top, otherwise it
    goes at the end of the page.
    """
    after: dict[int, list[tuple[BBox, int]]] = {}
    trailing: list[tuple[BBox, int]] = []
    for bbox, xref in images:
        above = [i for i, s in enumerate(segments)
                 if s.bbox.y1 <= bbox.y0 + 2 and min(s.bbox.x1, bbox.x1) - max(s.bbox.x0, bbox.x0) > 0]
        if above:
            idx = max(above, key=lambda i: segments[i].bbox.y1)
            after.setdefault(idx, []).append((bbox, xref))
            continue
        below = [i for i, s in enumerate(segments)
                 if s.bbox.y0 >= bbox.y1 - 2 and min(s.bbox.x1, bbox.x1) - max(s.bbox.x0, bbox.x0) > 0]
        if not below:
            below = [i for i, s in enumerate(segments) if s.bbox.y0 >= bbox.y0]
        if below:
            after.setdefault(min(below) - 1, []).append((bbox, xref))
        else:
            trailing.append((bbox, xref))
    return after, trailing


def _emit_page_from_segments(document: docx.document.Document, pdf: pymupdf.Document, page_index: int,
                             segments: list[TextSegment]) -> int:
    page = pdf[page_index]
    after, trailing = _image_slots(segments, _page_images(page))
    pictures = 0

    def emit_images(items: list[tuple[BBox, int]]) -> None:
        nonlocal pictures
        for bbox, xref in items:
            data = _image_bytes(pdf, xref)
            if data is not None and _add_picture(document, data, bbox):
                pictures += 1

    emit_images(after.get(-1, []))
    for idx, seg in enumerate(segments):
        _add_segment(document, seg)
        emit_images(after.get(idx, []))
    emit_images(trailing)
    return pictures


def _block_text(raw: str) -> str:
    lines = [ln.strip() for ln in raw.splitlines() if ln.strip()]
    if not lines:
        return ""
    lang = detect_language(" ".join(lines))
    joiner = "" if lang is not None and is_cjk(lang) else " "
    return joiner.join(lines)


def _emit_page_from_blocks(document: docx.document.Document, pdf: pymupdf.Document, page_index: int) -> int:
    """No segment information: paragraphs from ``get_text("blocks")`` in content order."""
    page = pdf[page_index]
    images = _page_images(page)
    pictures = 0
    # image blocks included; ligature glyphs (ﬁ, ﬂ) expanded to plain letters for Word
    flags = (pymupdf.TEXTFLAGS_BLOCKS | pymupdf.TEXT_PRESERVE_IMAGES) & ~pymupdf.TEXT_PRESERVE_LIGATURES
    for x0, y0, x1, y1, text, _no, btype in page.get_text("blocks", flags=flags):
        bbox = BBox(x0=x0, y0=y0, x1=x1, y1=y1)
        if btype == 0:
            para = _block_text(text)
            if para:
                document.add_paragraph(para)
            continue
        match = max(images, key=lambda im: im[0].intersection_area(bbox), default=None)
        data: Optional[bytes] = None
        if match is not None and match[0].intersection_area(bbox) >= 0.5 * max(bbox.area, 1e-6):
            data = _image_bytes(pdf, match[1])
        if data is None:
            clip = bbox.to_rect() & page.rect
            if clip.is_empty:
                continue
            data = page.get_pixmap(dpi=150, clip=clip).tobytes("png")
        if _add_picture(document, data, bbox):
            pictures += 1
    return pictures


def _display_title(doc: TranslatedDocument) -> str:
    """The title in the target language: ``doc.title`` is recorded by the extractor
    from the *source* text, so the translation of the matching segment is used
    when there is one."""
    title = doc.title.strip()
    if not title:
        return ""
    for seg in doc.segments:
        if seg.kind == SegmentKind.TEXT and seg.source_text.strip() == title:
            return seg.effective_text.strip() or title
    return title


def build_docx_from_segments(translated_pdf: PathLike, out_path: PathLike,
                             doc: Optional[TranslatedDocument] = None) -> str:
    """Build a DOCX with python-docx: headings/paragraphs from ``doc.segments`` in
    ``(page, reading_order)`` order with the page images at their reading
    position, or from ``page.get_text("blocks")`` when ``doc`` is None.

    The document title (translated) goes into the core properties; a "Title"
    paragraph is added only when no segment already carries that text, so the
    first heading is not duplicated."""
    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    pdf = pymupdf.open(str(translated_pdf))
    try:
        document = docx.Document()
        _configure_fonts(document, doc.target_lang if doc is not None else None)
        title = _display_title(doc) if doc is not None else ""
        if title:
            document.core_properties.title = title
            if not any(s.kind == SegmentKind.TEXT and s.effective_text.strip() == title for s in doc.segments):
                document.add_heading(title, level=0)
        by_page: dict[int, list[TextSegment]] = {}
        if doc is not None:
            for seg in doc.segments:
                if seg.kind != SegmentKind.TEXT:
                    continue
                if not 0 <= seg.page < pdf.page_count:
                    log.warning("segment %s refers to page %d but the PDF has %d page(s); left out of the DOCX",
                                seg.id, seg.page, pdf.page_count)
                    continue
                by_page.setdefault(seg.page, []).append(seg)
        pictures = 0
        for pno in range(pdf.page_count):
            if pno > 0:
                document.add_page_break()
            if doc is not None:
                segs = sorted(by_page.get(pno, []), key=lambda s: (s.reading_order, s.bbox.y0, s.bbox.x0))
                pictures += _emit_page_from_segments(document, pdf, pno, segs)
            else:
                pictures += _emit_page_from_blocks(document, pdf, pno)
        paragraphs = len(document.paragraphs)
        document.save(str(out))
    finally:
        pdf.close()
    log.info("python-docx wrote %s (%d paragraphs, %d pictures)", out, paragraphs, pictures)
    return str(out)


# --------------------------------------------------------------------------- #
# Entry point
# --------------------------------------------------------------------------- #


def export_docx(translated_pdf: PathLike, out_path: PathLike, doc: Optional[TranslatedDocument] = None,
                timeout: float = 180, *, use_libreoffice: Optional[bool] = None) -> str:
    """Write a DOCX version of ``translated_pdf`` and return its path.

    ``timeout`` (seconds) bounds the LibreOffice conversion. ``use_libreoffice``
    (keyword-only): ``None`` (default) tries LibreOffice when it is installed
    and falls back to python-docx when the conversion fails or times out;
    ``True`` requires LibreOffice and raises :class:`ExportError` when it
    cannot deliver; ``False`` always uses the python-docx builder.
    """
    if use_libreoffice is False:
        return build_docx_from_segments(translated_pdf, out_path, doc)
    exe = libreoffice_executable()
    if exe is None:
        if use_libreoffice:
            raise ExportError("LibreOffice (soffice) is not installed or not on PATH")
        log.info("LibreOffice not found; building the DOCX with python-docx")
        return build_docx_from_segments(translated_pdf, out_path, doc)
    try:
        return convert_pdf_to_docx_libreoffice(translated_pdf, out_path, timeout=timeout)
    except ExportError as exc:
        if use_libreoffice:
            raise
        log.warning("LibreOffice conversion failed (%s); building the DOCX with python-docx", exc)
        return build_docx_from_segments(translated_pdf, out_path, doc)
